"""Isolated synthetic records only; never writes to Drive or live execution.json."""
import copy
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
import sync_execution as sync

TODAY=date(2026,10,26)


def record(state='测试中'):
    r=[None]*26
    r[:13]=['ISOLATED-TEST-001','TEST:US','TEST-PARENT','TEST-CHILD','TEST-MSKU','TEST-CANDIDATE',10,10.3,10.3,'USD',date(2026,10,19),None,state]
    r[19]='https://example.invalid/isolated-test'
    return r


def finished():
    r=record('已完成');r[11]=date(2026,10,25)
    r[14:19]=[100,125,'是',999999,'保留'];r[23]='通过'
    return r


def queue():
    r=[None]*31;r[1:5]=['TEST:US','TEST-PARENT','TEST-CHILD','TEST-MSKU']
    r[6]=10;r[12]='符合';r[15:21]=['通过']*6
    r[21:26]=['operator',TODAY,'evidence',TODAY,date(2026,10,31)];r[29]=10
    return r


def inputs(rows,qs=None):
    return {'李漉':{'rows':rows,'queue':qs or [],'asof':str(TODAY)}}


class ExecutionTests(unittest.TestCase):
    def test_parent_name_column_preserves_canonical_input_and_states(self):
        import io
        from openpyxl import Workbook
        config={'sheet':'02｜提价执行与复盘','queue_sheet':'01｜本周待提价清单'}
        parsed=[]
        for extra in (False,True):
            w=Workbook();w.remove(w.active);ex=w.create_sheet(config['sheet']);q=w.create_sheet(config['queue_sheet'])
            for j,h in enumerate(sync.HEADERS,1):ex.cell(7,j,h)
            for j,v in enumerate(record(),1):ex.cell(8,j,v)
            h=['field'+str(i) for i in range(31)]
            for j,v in zip((2,3,4,5,13,30),('站点','父ASIN','子ASIN','执行MSKU','候选资格','已核现价原币')):h[j-1]=v
            r=queue()
            if extra:h.insert(3,'父单名');r.insert(3,'Existing explicit parent label')
            for j,v in enumerate(h,1):q.cell(7,j,v)
            for j,v in enumerate(r,1):q.cell(8,j,v)
            q.cell(3,2,TODAY)
            output=io.BytesIO();w.save(output)
            parsed.append(sync.read_workbook(output.getvalue(),config))
        self.assertEqual(parsed[0],parsed[1])
        self.assertEqual(sync.summaries({'李漉':parsed[0]},TODAY),sync.summaries({'李漉':parsed[1]},TODAY))

    def metrics(self,rows,qs=None):
        return sync.summaries(inputs(rows,qs),TODAY)['李漉']['operational']

    def test_blank(self):
        m=self.metrics([]);self.assertEqual(m['active_count'],0);self.assertIsNone(m['verified_increment_cny'])

    def test_incomplete_status_not_actual(self):
        r=record();r[19]=None;m=self.metrics([r]);self.assertEqual(m['active_count'],0);self.assertEqual(m['registered_occupancy_count'],1)

    def test_started(self):
        m=self.metrics([record()]);self.assertEqual(m['active_count'],1);self.assertEqual(m['completed_count'],0)

    def test_future_start(self):
        r=record();r[10]=date(2026,10,27);self.assertFalse(sync.started(r,TODAY))

    def test_invalid_completion(self):
        r=record('已完成');self.assertEqual(self.metrics([r])['completed_count'],0)

    def test_completed_not_verified_without_comparison(self):
        r=finished();r[16]='未确认';m=self.metrics([r]);self.assertEqual(m['completed_count'],1);self.assertIsNone(m['verified_increment_cny'])

    def test_comparable_profit_uses_raw_parent_difference_once(self):
        m=self.metrics([finished()]);self.assertEqual(m['verified_increment_cny'],25);self.assertEqual(m['success_rate'],1)

    def test_missing_fx_proof(self):
        r=finished();r[23]='待核验';self.assertIsNone(self.metrics([r])['verified_increment_cny'])

    def test_bad_end(self):
        r=finished();r[11]=date(2026,10,18);self.assertFalse(sync.completed(r,TODAY))

    def test_duplicate_same_record(self):
        rows=sync.dedupe([finished(),finished()]);self.assertEqual(self.metrics(rows)['verified_increment_cny'],25)

    def test_conflicting_test_id(self):
        r=finished();x=copy.deepcopy(r);x[15]=130
        with self.assertRaises(ValueError):sync.dedupe([r,x])

    def test_parent_overlap(self):
        a=finished();b=finished();b[0]='OTHER';b[3]='OTHER-CHILD'
        self.assertIsNone(self.metrics([a,b])['verified_increment_cny'])

    def test_cross_manager_parent_overlap(self):
        i=inputs([finished()]);r=finished();r[0]='OTHER';r[3]='OTHER'
        i['陈梦怡']={'rows':[r],'queue':[],'asof':str(TODAY)}
        result=sync.summaries(i,TODAY)
        self.assertTrue(all(g['operational']['verified_increment_cny'] is None for g in result.values()))

    def test_actual_testing_blocks_ready(self):
        m=self.metrics([record()],[queue()]);self.assertEqual(m['ready_count'],0)
        self.assertEqual(m['execution_states']['已在测试（不新建）'],1)

    def test_complete_checks_ready(self):
        self.assertEqual(self.metrics([],[queue()])['ready_count'],1)

    def test_missing_check_not_ready(self):
        r=queue();r[15]='待核验';m=self.metrics([],[r]);self.assertEqual(m['ready_count'],0);self.assertEqual(m['pending_verification'],1)

    def test_status_only_never_completed(self):
        for state in ['填写中','待核验','',None,'已完成']:
            r=record(state);r[19]=None
            self.assertEqual(self.metrics([r])['completed_count'],0)

    def test_no_theoretical_profit_override(self):
        self.assertNotIn('theoretical_increment_cny',self.metrics([finished()]))

    def test_site_isolation(self):
        result=sync.summaries(inputs([finished()]),TODAY)['李漉']
        self.assertEqual(result['site_operational']['TEST:US']['verified_increment_cny'],25)

    def test_failed_cloud_read_retains_last_good(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);config={'sources':[{'manager':'李漉','file_id':'test'}]}
            (root/'execution_sources.json').write_text(json.dumps(config),encoding='utf8')
            old={'schema':'execution-sync-v1','groups':{'李漉':{'operational':{'active_count':7},'last_success_at':'before'}}}
            sync.write_outputs(root,old)
            with patch.object(sync,'DriveReader') as reader:
                reader.return_value.metadata.side_effect=ValueError('drive_metadata_http_403')
                with patch.object(sync,'cache_key',return_value=b'0'*32):
                    result=sync.run(root)
            self.assertEqual(result['groups']['李漉']['operational']['active_count'],7)
            self.assertEqual(result['groups']['李漉']['last_success_at'],'before')


if __name__=='__main__':unittest.main()
